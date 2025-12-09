import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from database import get_db, Employee, TimeRecord, Transaction
from config import EXPORT_DIR
import os
from datetime import datetime
from typing import List, Dict, Optional

class ExportManager:
    def __init__(self):
        self.export_dir = EXPORT_DIR
        os.makedirs(self.export_dir, exist_ok=True)
    
    def export_employees_to_excel(self, filename: str = None) -> str:
        """Export all employees to Excel file"""
        if not filename:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"employees_{timestamp}.xlsx"
        
        filepath = os.path.join(self.export_dir, filename)
        
        with get_db() as db:
            employees = db.query(Employee).all()
            
            # Create DataFrame
            data = []
            for emp in employees:
                data.append({
                    'ID': emp.id,
                    'Telegram ID': emp.telegram_id,
                    'ФИО': emp.full_name,
                    'Должность': emp.position,
                    'Ставка в час': emp.hourly_rate,
                    'Баланс': emp.balance,
                    'Роль': emp.role,
                    'Дата регистрации': emp.created_at
                })
            
            df = pd.DataFrame(data)
            
            # Create workbook and worksheet
            wb = Workbook()
            ws = wb.active
            ws.title = "Сотрудники"
            
            # Add data to worksheet
            for r in dataframe_to_rows(df, index=False, header=True):
                ws.append(r)
            
            # Format headers
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # Auto-adjust column widths
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width
            
            wb.save(filepath)
            return filepath
    
    def export_timesheets_to_excel(self, start_date: datetime = None, end_date: datetime = None, filename: str = None) -> str:
        """Export timesheets to Excel file with optional date filtering"""
        if not filename:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"timesheets_{timestamp}.xlsx"
        
        filepath = os.path.join(self.export_dir, filename)
        
        with get_db() as db:
            # Build query with optional date filtering
            query = db.query(TimeRecord).join(Employee)
            
            if start_date:
                query = query.filter(TimeRecord.date >= start_date)
            if end_date:
                query = query.filter(TimeRecord.date <= end_date)
            
            time_records = query.all()
            
            # Create DataFrame
            data = []
            for record in time_records:
                data.append({
                    'ID табеля': record.id,
                    'ID сотрудника': record.employee_id,
                    'ФИО': record.employee.full_name,
                    'Отработано часов': record.hours_worked,
                    'Дата': record.date,
                    'Статус': record.status,
                    'Комментарий': record.comment or ''
                })
            
            df = pd.DataFrame(data)
            
            # Create workbook and worksheet
            wb = Workbook()
            ws = wb.active
            ws.title = "Табели"
            
            # Add data to worksheet
            for r in dataframe_to_rows(df, index=False, header=True):
                ws.append(r)
            
            # Format headers
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
            
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # Auto-adjust column widths
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width
            
            wb.save(filepath)
            return filepath
    
    def export_transactions_to_excel(self, start_date: datetime = None, end_date: datetime = None, filename: str = None) -> str:
        """Export transactions (bonuses/penalties) to Excel file"""
        if not filename:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"transactions_{timestamp}.xlsx"
        
        filepath = os.path.join(self.export_dir, filename)
        
        with get_db() as db:
            # Build query with optional date filtering
            query = db.query(Transaction).join(Employee)
            
            if start_date:
                query = query.filter(Transaction.date >= start_date)
            if end_date:
                query = query.filter(Transaction.date <= end_date)
            
            transactions = query.all()
            
            # Create DataFrame
            data = []
            for trans in transactions:
                data.append({
                    'ID транзакции': trans.id,
                    'ID сотрудника': trans.employee_id,
                    'ФИО': trans.employee.full_name,
                    'Сумма': trans.amount,
                    'Тип': trans.transaction_type,
                    'Дата': trans.date,
                    'Обработано админом': trans.processed_by,
                    'Комментарий': trans.comment or ''
                })
            
            df = pd.DataFrame(data)
            
            # Create workbook and worksheet
            wb = Workbook()
            ws = wb.active
            ws.title = "Транзакции"
            
            # Add data to worksheet
            for r in dataframe_to_rows(df, index=False, header=True):
                ws.append(r)
            
            # Format headers
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="FF9900", end_color="FF9900", fill_type="solid")
            
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # Auto-adjust column widths
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width
            
            wb.save(filepath)
            return filepath
    
    def export_salary_report(self, month: int = None, year: int = None, filename: str = None) -> str:
        """Export salary report for a specific month/year"""
        if not filename:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            if month and year:
                filename = f"salary_report_{year}_{month:02d}_{timestamp}.xlsx"
            else:
                filename = f"salary_report_{timestamp}.xlsx"
        
        filepath = os.path.join(self.export_dir, filename)
        
        with get_db() as db:
            employees = db.query(Employee).all()
            
            # Calculate salary data for each employee
            data = []
            for emp in employees:
                # Calculate total hours for the specified month/year if provided
                time_records_query = db.query(TimeRecord).filter(TimeRecord.employee_id == emp.id)
                
                if month and year:
                    time_records_query = time_records_query.filter(
                        TimeRecord.date.month == month,
                        TimeRecord.date.year == year
                    )
                
                time_records = time_records_query.all()
                total_hours = sum(tr.hours_worked for tr in time_records if tr.status == 'approved')
                
                # Calculate bonuses and penalties for the specified month/year if provided
                transactions_query = db.query(Transaction).filter(Transaction.employee_id == emp.id)
                
                if month and year:
                    transactions_query = transactions_query.filter(
                        Transaction.date.month == month,
                        Transaction.date.year == year
                    )
                
                transactions = transactions_query.all()
                total_bonus_penalty = sum(t.amount for t in transactions)
                
                # Calculate total salary
                total_salary = (total_hours * emp.hourly_rate) + total_bonus_penalty
                
                data.append({
                    'ID': emp.id,
                    'ФИО': emp.full_name,
                    'Должность': emp.position,
                    'Отработано часов': total_hours,
                    'Почасовая ставка': emp.hourly_rate,
                    'Начислено за часы': total_hours * emp.hourly_rate,
                    'Премии/штрафы': total_bonus_penalty,
                    'Итого к выплате': total_salary,
                    'Текущий баланс': emp.balance
                })
            
            df = pd.DataFrame(data)
            
            # Create workbook and worksheet
            wb = Workbook()
            ws = wb.active
            ws.title = "Зарплаты"
            
            # Add data to worksheet
            for r in dataframe_to_rows(df, index=False, header=True):
                ws.append(r)
            
            # Format headers
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="5B9BD5", end_color="5B9BD5", fill_type="solid")
            
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # Format currency columns
            currency_cols = ['Почасовая ставка', 'Начислено за часы', 'Премии/штрафы', 'Итого к выплате', 'Текущий баланс']
            currency_col_indices = []
            for i, cell in enumerate(ws[1]):
                if cell.value in currency_cols:
                    currency_col_indices.append(i + 1)
            
            for row in ws.iter_rows(min_row=2):
                for col_idx in currency_col_indices:
                    if col_idx <= len(row):
                        cell = row[col_idx - 1]
                        cell.number_format = '#,##0.00'
            
            # Auto-adjust column widths
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width
            
            wb.save(filepath)
            return filepath
    
    def export_to_csv(self, query_result, columns, filename: str) -> str:
        """Generic method to export query results to CSV"""
        if not filename.endswith('.csv'):
            filename += '.csv'
        
        filepath = os.path.join(self.export_dir, filename)
        
        # Convert query result to DataFrame
        data = []
        for row in query_result:
            row_data = {}
            for i, col in enumerate(columns):
                row_data[col] = getattr(row, col) if hasattr(row, col) else row[i] if isinstance(row, (list, tuple)) else row
            data.append(row_data)
        
        df = pd.DataFrame(data)
        df.to_csv(filepath, index=False, encoding='utf-8-sig')
        
        return filepath
    
    def export_multiple_sheets(self, filename: str = None) -> str:
        """Export multiple sheets to a single Excel file"""
        if not filename:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"complete_report_{timestamp}.xlsx"
        
        filepath = os.path.join(self.export_dir, filename)
        
        with get_db() as db:
            # Create workbook
            wb = Workbook()
            
            # Remove default sheet
            wb.remove(wb.active)
            
            # Add employees sheet
            employees_ws = wb.create_sheet("Сотрудники")
            employees = db.query(Employee).all()
            emp_data = []
            for emp in employees:
                emp_data.append({
                    'ID': emp.id,
                    'Telegram ID': emp.telegram_id,
                    'ФИО': emp.full_name,
                    'Должность': emp.position,
                    'Ставка в час': emp.hourly_rate,
                    'Баланс': emp.balance,
                    'Роль': emp.role
                })
            
            emp_df = pd.DataFrame(emp_data)
            for r in dataframe_to_rows(emp_df, index=False, header=True):
                employees_ws.append(r)
            
            # Format employees sheet
            self._format_sheet(employees_ws, "366092")
            
            # Add timesheets sheet
            timesheets_ws = wb.create_sheet("Табели")
            time_records = db.query(TimeRecord).join(Employee).all()
            ts_data = []
            for record in time_records:
                ts_data.append({
                    'ID табеля': record.id,
                    'ФИО': record.employee.full_name,
                    'Отработано часов': record.hours_worked,
                    'Дата': record.date,
                    'Статус': record.status
                })
            
            ts_df = pd.DataFrame(ts_data)
            for r in dataframe_to_rows(ts_df, index=False, header=True):
                timesheets_ws.append(r)
            
            # Format timesheets sheet
            self._format_sheet(timesheets_ws, "70AD47")
            
            # Add transactions sheet
            transactions_ws = wb.create_sheet("Транзакции")
            transactions = db.query(Transaction).join(Employee).all()
            trans_data = []
            for trans in transactions:
                trans_data.append({
                    'ID транзакции': trans.id,
                    'ФИО': trans.employee.full_name,
                    'Сумма': trans.amount,
                    'Тип': trans.transaction_type,
                    'Дата': trans.date,
                    'Комментарий': trans.comment or ''
                })
            
            trans_df = pd.DataFrame(trans_data)
            for r in dataframe_to_rows(trans_df, index=False, header=True):
                transactions_ws.append(r)
            
            # Format transactions sheet
            self._format_sheet(transactions_ws, "FF9900")
            
            wb.save(filepath)
            return filepath
    
    def _format_sheet(self, ws, color_code: str):
        """Helper method to format a worksheet"""
        # Format headers
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color=color_code, end_color=color_code, fill_type="solid")
        
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

# Create a singleton instance
export_manager = ExportManager()